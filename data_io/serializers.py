"""Importación y exportación de problemas y resultados (Excel / CSV).

Formato de la tabla del problema (una hoja de Excel o un CSV)::

    Alternativa | Precio | Calidad   | ...
    Proveedor A |    120 |         8 | ...
    Proveedor B |    100 |         7 | ...
    Sentido     |  Costo | Beneficio | ...   (opcional: Max/Min, Beneficio/Costo)
    Peso        |   0,35 |      0,30 | ...   (opcional)

* La primera fila no vacía es el encabezado (nombres de los criterios).
* La primera columna contiene los nombres de las alternativas.
* Las filas «Sentido» y «Peso» pueden ir antes o después de los datos. Una fila
  sin rótulo cuyos valores son todos MAX/MIN también se reconoce como sentido.
* Se aceptan coma o punto decimal y CSV separado por «,» o «;».
* La lectura se detiene en la primera fila vacía después de los datos, de modo
  que se pueden importar planillas con otras tablas debajo.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field

import numpy as np
import pandas as pd

from .validators import BENEFIT, COST, is_blank, parse_direction, parse_number, unique_labels

_DIRECTION_LABELS = {"sentido", "sentidos", "tipo", "direccion", "dirección", "objetivo", "max/min", "optimizacion", "optimización"}
_WEIGHT_LABELS = {"peso", "pesos", "w", "wj", "w_j", "ponderacion", "ponderación", "ponderaciones", "importancia"}

EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass
class ImportedProblem:
    alternatives: list[str]
    criteria: list[str]
    values: pd.DataFrame  # m×n, float con NaN en celdas vacías o inválidas
    directions: list[str] | None = None  # "Beneficio" / "Costo"
    weights: list[float | None] | None = None
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Lectura
# ---------------------------------------------------------------------------

def _read_raw(data: bytes, filename: str, sheet: str | int = 0) -> pd.DataFrame:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        return pd.read_excel(io.BytesIO(data), sheet_name=sheet, header=None, dtype=object)
    if name.endswith((".csv", ".txt")):
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("No se pudo leer la codificación del archivo CSV.")
        first = next((line for line in text.splitlines() if line.strip()), "")
        sep = ";" if first.count(";") > first.count(",") else ("\t" if "\t" in first else ",")
        return pd.read_csv(io.StringIO(text), sep=sep, header=None, dtype=str, keep_default_na=False, skip_blank_lines=False)
    raise ValueError("Formato no soportado: use un archivo .xlsx, .xls o .csv.")


def excel_sheet_names(data: bytes) -> list[str]:
    return pd.ExcelFile(io.BytesIO(data)).sheet_names


def read_problem(data: bytes, filename: str, sheet: str | int = 0) -> ImportedProblem:
    """Lee un problema desde bytes de un archivo Excel o CSV."""
    try:
        raw = _read_raw(data, filename, sheet)
    except ValueError:
        raise
    except Exception as exc:  # errores de formato de pandas/openpyxl
        raise ValueError(f"No se pudo leer el archivo: {exc}") from exc
    return parse_problem_table(raw)


def _is_direction_row(cells: list) -> bool:
    values = [v for v in cells if not is_blank(v)]
    return bool(values) and all(parse_direction(v) is not None for v in values)


def parse_problem_table(raw: pd.DataFrame) -> ImportedProblem:
    """Interpreta una tabla cruda (sin encabezado) como problema MOORA."""
    df = raw.copy()
    df = df.map(lambda v: None if is_blank(v) else v)
    df = df.dropna(axis=1, how="all")
    if df.empty:
        raise ValueError("El archivo está vacío.")
    rows = df.to_numpy(dtype=object).tolist()

    warnings: list[str] = []
    header: list | None = None
    cols: list[int] = []
    directions_row: list | None = None
    weights_row: list | None = None
    data_rows: list[list] = []
    ended = False  # se encontró una fila vacía después de los datos

    for row in rows:
        label = row[0]
        label_key = "" if is_blank(label) else str(label).strip().lower()
        rest = row[1:]
        in_table = [rest[c - 1] for c in cols] if header is not None else rest

        if all(is_blank(v) for v in [label, *in_table]):
            ended = ended or bool(data_rows)
            continue
        if label_key in _WEIGHT_LABELS:
            weights_row = in_table if header is not None else rest
            continue
        if label_key in _DIRECTION_LABELS or (not label_key and _is_direction_row(in_table)):
            directions_row = in_table if header is not None else rest
            continue
        if ended:
            break  # después de la tabla solo se aceptan filas de sentido o peso
        if header is None:
            header = row
            cols = [j for j in range(1, len(row)) if not is_blank(row[j])]
            if not cols:
                raise ValueError("No se encontraron nombres de criterios en el encabezado.")
            # Si sentido/peso aparecieron antes del encabezado, recortarlos a las columnas de la tabla.
            if directions_row is not None and len(directions_row) == len(rest):
                directions_row = [directions_row[c - 1] for c in cols]
            if weights_row is not None and len(weights_row) == len(rest):
                weights_row = [weights_row[c - 1] for c in cols]
            continue
        if all(is_blank(v) for v in in_table):
            warnings.append(f"Se ignoró la fila «{label}» porque no tiene valores.")
            continue
        data_rows.append([label, *in_table])

    if header is None:
        raise ValueError("No se encontró el encabezado con los nombres de los criterios.")
    if not data_rows:
        raise ValueError("No se encontraron filas de alternativas debajo del encabezado.")

    criteria = unique_labels([header[c] for c in cols], "C")
    alternatives = unique_labels([r[0] for r in data_rows], "A")
    if [str(header[c]).strip() for c in cols] != criteria:
        warnings.append("Se completaron o renombraron criterios sin nombre o repetidos.")
    if [("" if is_blank(r[0]) else str(r[0]).strip()) for r in data_rows] != alternatives:
        warnings.append("Se completaron o renombraron alternativas sin nombre o repetidas.")

    values = np.full((len(data_rows), len(cols)), np.nan)
    invalid = []
    for i, r in enumerate(data_rows):
        for j, v in enumerate(r[1:]):
            try:
                f = parse_number(v)
            except ValueError:
                invalid.append(f"{alternatives[i]} / {criteria[j]} = «{v}»")
                continue
            if f is not None:
                values[i, j] = f
    if invalid:
        listed = ", ".join(invalid[:8]) + (f" y {len(invalid) - 8} más" if len(invalid) > 8 else "")
        warnings.append(f"Valores no numéricos que quedaron vacíos: {listed}.")

    directions = None
    if directions_row is not None:
        parsed = [parse_direction(v) for v in directions_row]
        directions = [BENEFIT if p else COST if p is False else BENEFIT for p in parsed]
        unknown = [criteria[j] for j, p in enumerate(parsed) if p is None]
        if unknown:
            warnings.append(f"Sentido no reconocido en {', '.join(unknown)}: se asumió Beneficio.")

    weights = None
    if weights_row is not None:
        weights = []
        for j, v in enumerate(weights_row):
            try:
                weights.append(parse_number(v))
            except ValueError:
                weights.append(None)
                warnings.append(f"Peso no numérico en {criteria[j]}: «{v}».")

    return ImportedProblem(
        alternatives=alternatives,
        criteria=criteria,
        values=pd.DataFrame(values, columns=criteria, index=alternatives),
        directions=directions,
        weights=weights,
        warnings=warnings,
    )


# ---------------------------------------------------------------------------
# Escritura
# ---------------------------------------------------------------------------

def problem_frame(alternatives, criteria, values, directions, weights) -> pd.DataFrame:
    """Tabla del problema en el mismo formato que se importa (con filas Sentido y Peso)."""
    vals = np.asarray(values, dtype=object)
    rows = [[a, *vals[i].tolist()] for i, a in enumerate(alternatives)]
    rows.append(["Sentido", *directions])
    rows.append(["Peso", *weights])
    return pd.DataFrame(rows, columns=["Alternativa", *criteria])


def _autosize(ws) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 45)


def _write_sheets(sheets: dict[str, pd.DataFrame], index_flags: dict[str, bool]) -> bytes:
    from openpyxl.styles import Font, PatternFill

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name[:31], index=index_flags.get(name, True))
            ws = writer.sheets[name[:31]]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F3A5F")
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, float):
                        cell.number_format = "0.000000"
            ws.freeze_panes = "B2"
            _autosize(ws)
    return buffer.getvalue()


def problem_to_excel(alternatives, criteria, values, directions, weights) -> bytes:
    frame = problem_frame(alternatives, criteria, values, directions, weights)
    return _write_sheets({"Problema": frame}, {"Problema": False})


def problem_to_csv(alternatives, criteria, values, directions, weights, sep: str = ",", decimal: str = ".") -> bytes:
    frame = problem_frame(alternatives, criteria, values, directions, weights)
    return frame.to_csv(index=False, sep=sep, decimal=decimal).encode("utf-8-sig")


def result_sheets(result, comparison=None, sensitivity=None, entropy=None, problem_name: str = "") -> dict[str, pd.DataFrame]:
    """Tablas de todos los pasos intermedios y resultados, listas para exportar."""
    sheets: dict[str, pd.DataFrame] = {}
    sheets["Problema"] = problem_frame(
        result.alternatives, result.criteria, result.matrix, result.directions, result.weights
    ).set_index("Alternativa")
    sheets["Criterios"] = result.criteria_frame()
    if entropy is not None:
        sheets["Pesos entropia"] = pd.DataFrame(
            {"E_j (entropía)": entropy.entropy, "d_j = 1 - E_j": entropy.divergence, "w_j": entropy.weights},
            index=pd.Index(result.criteria, name="Criterio"),
        )
    sheets["Normalizada"] = result.normalized_frame()
    sheets["Ponderada"] = result.weighted_frame()
    sheets["Sistema de Razones"] = result.ratio_frame()
    sheets["Punto de referencia"] = result.reference_frame()
    sheets["Desviaciones"] = result.deviations_frame()
    sheets["Desviaciones ponderadas"] = result.weighted_deviations_frame()
    sheets["Resultados"] = result.results_frame()
    if comparison is not None:
        comp = comparison.table.copy()
        comp["Coincide"] = comp["Coincide"].map({True: "Sí", False: "No"})
        sheets["Comparacion"] = comp
        summary = {
            "Problema": problem_name or "-",
            "Mejor (Sistema de Razones)": " / ".join(comparison.ratio_winners),
            "Mejor (Punto de Referencia)": " / ".join(comparison.reference_winners),
            "Coincide la mejor alternativa": "Sí" if comparison.winners_match else "No",
            "Spearman rho": comparison.spearman,
            "Posiciones coincidentes": f"{comparison.position_matches} de {len(comparison.table)}",
        }
        sheets["Resumen"] = pd.DataFrame({"Valor": list(summary.values())}, index=pd.Index(summary, name="Indicador"))
    if sensitivity is not None:
        sheets["Sensibilidad"] = sensitivity.winners_frame()
        sheets["Sensibilidad detalle"] = sensitivity.long_frame()
        if not sensitivity.change_points.empty:
            sheets["Puntos de cambio"] = sensitivity.change_points
    if "Resumen" in sheets:  # el resumen va primero
        sheets = {"Resumen": sheets.pop("Resumen"), **sheets}
    return sheets


_NO_INDEX = {"Sensibilidad", "Sensibilidad detalle", "Puntos de cambio"}


def results_to_excel(result, comparison=None, sensitivity=None, entropy=None, problem_name: str = "") -> bytes:
    """Libro Excel con una hoja por paso intermedio."""
    sheets = result_sheets(result, comparison, sensitivity, entropy, problem_name)
    return _write_sheets(sheets, {k: k not in _NO_INDEX for k in sheets})


def results_to_csv_zip(result, comparison=None, sensitivity=None, entropy=None, problem_name: str = "",
                       sep: str = ",", decimal: str = ".") -> bytes:
    """ZIP con un CSV por tabla."""
    sheets = result_sheets(result, comparison, sensitivity, entropy, problem_name)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for i, (name, frame) in enumerate(sheets.items(), start=1):
            fname = f"{i:02d}_{name.replace(' ', '_').lower()}.csv"
            text = frame.to_csv(index=name not in _NO_INDEX, sep=sep, decimal=decimal)
            zf.writestr(fname, text.encode("utf-8-sig"))
    return buffer.getvalue()


def results_to_csv(result, comparison=None, sep: str = ",", decimal: str = ".") -> bytes:
    """CSV único con los resultados finales de ambos métodos."""
    frame = comparison.table.copy() if comparison is not None else result.results_frame()
    if "Coincide" in frame:
        frame["Coincide"] = frame["Coincide"].map({True: "Sí", False: "No"})
    return frame.to_csv(sep=sep, decimal=decimal).encode("utf-8-sig")
